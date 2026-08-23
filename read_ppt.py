"""Read PPT template structure."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from pptx import Presentation
from pptx.util import Inches, Pt, Emu

prs = Presentation(r'C:\D_Games\capstone\ppts\sem 6 esa\UE23CS320B_Phase2 -ESA.pptx')

print(f"Slide width: {prs.slide_width}, height: {prs.slide_height}")
print(f"Total slides: {len(prs.slides)}")
print(f"Slide layouts: {[l.name for l in prs.slide_layouts]}")

print("\n" + "="*80)
for i, slide in enumerate(prs.slides):
    print(f"\n--- SLIDE {i+1} ---")
    print(f"  Layout: {slide.slide_layout.name}")
    for j, shape in enumerate(slide.shapes):
        print(f"  Shape {j}: {shape.shape_type}, name='{shape.name}'")
        print(f"    Pos: left={shape.left}, top={shape.top}, w={shape.width}, h={shape.height}")
        if shape.has_text_frame:
            for k, para in enumerate(shape.text_frame.paragraphs):
                text = para.text.strip()
                if text:
                    font_info = ""
                    if para.runs:
                        r = para.runs[0]
                        font_info = f" [font={r.font.name}, size={r.font.size}, bold={r.font.bold}]"
                    print(f"    P{k}: '{text[:100]}'{font_info}")
