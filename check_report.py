"""Extract image locations from the docx report."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
from docx.oxml.ns import qn

d = Document(r'C:\D_Games\capstone\Report\sem 6\CAPSTONE PHASE 2 REPORT_TODAY.docx')

print("=== IMAGES IN DOCUMENT ===")
for i, p in enumerate(d.paragraphs):
    has_img = False
    for r in p.runs:
        drawings = r.element.findall(qn('w:drawing'))
        picts = r.element.findall(qn('w:pict'))
        if drawings or picts:
            has_img = True
    if has_img:
        txt = p.text.strip()[:80] if p.text.strip() else "[image only]"
        print(f"  Para {i}: {txt}")

print("\n=== ALL TABLES ===")
for i, tbl in enumerate(d.tables):
    first_row = []
    if tbl.rows:
        for cell in tbl.rows[0].cells:
            first_row.append(cell.text.strip()[:20])
    print(f"  Table {i+1}: {' | '.join(first_row)}")

print("\n=== FIGURE CAPTIONS (text containing 'Fig' or 'figure') ===")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t and ('Fig' in t or 'figure' in t.lower()):
        print(f"  Para {i}: {t[:120]}")

print(f"\n=== TOTAL: {len(d.tables)} tables, {sum(1 for r in d.part.rels.values() if 'image' in r.reltype)} images ===")
