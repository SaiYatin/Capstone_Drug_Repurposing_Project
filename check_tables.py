import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

d = Document(r'C:\D_Games\capstone\Report\sem 6\CAPSTONE PHASE 2 REPORT_TODAY.docx')

for i, t in enumerate(d.tables[7:], start=8):
    print(f"\n=== TABLE {i} ===")
    for r in t.rows:
        cells = [c.text.strip()[:25] for c in r.cells]
        print("  " + " | ".join(cells))

# Also check the results/conclusion text for disease mentions
print("\n\n=== DISEASES MENTIONED IN RESULTS ===")
for i, p in enumerate(d.paragraphs[444:497]):
    t = p.text.strip()
    if t:
        print(f"  {444+i}: {t[:200]}")
